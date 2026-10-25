"""Saved results and the Excel report.

Answers and judgements are appended to results/*.jsonl as they arrive. The Excel report is rebuilt from those
files after every run and updated in place, so it always shows every model and case evaluated so far.
"""

import json
import sys
import threading
from statistics import mean

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .llm import DIMENSIONS


# ---------- saved results ----------

class JsonlStore:
    """Append-only JSONL file keyed by record["key"]; the last record for a key wins. Thread-safe writes."""

    def __init__(self, path):
        self.path = path
        self.lock = threading.Lock()
        self.records = {}
        if path.exists():
            with open(path, encoding="utf-8") as f:
                for line in f:
                    if line.strip():
                        record = json.loads(line)
                        self.records[record["key"]] = record

    def __contains__(self, key):
        return key in self.records

    def add(self, record):
        with self.lock:
            self.records[record["key"]] = record
            with open(self.path, "a", encoding="utf-8") as f:
                f.write(json.dumps(record, ensure_ascii=False) + "\n")


def answer_key(item, model, run):
    return f"{item['case_id']}|{item['question_key']}|{model}|{run}"


def is_current(judgement, answers_by_key):
    """True when the judgement was made on the latest saved version of its answer."""
    return judgement.get("answer_created_at") == answers_by_key[judgement["answer_key"]].get("created_at")


# ---------- Excel report ----------

def clean(value):
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def avg(values):
    values = [v for v in values if v is not None]
    return round(mean(values), 2) if values else None


def add_sheet(wb, title, headers, rows, wrap=()):
    """Write a sheet, replacing the contents of an existing sheet of that name in place.

    Clearing the cells rather than deleting the sheet keeps other sheets' formulas that point at it working.
    """
    if title in wb.sheetnames:
        ws = wb[title]
        ws.auto_filter.ref = None
        if ws.max_row:
            ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append([clean(v) for v in row])
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        letter = get_column_letter(col)
        if header in wrap:
            ws.column_dimensions[letter].width = 60
            for r in range(2, len(rows) + 2):
                ws.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        else:
            width = max([len(str(header))] + [len(str(r[col - 1])) for r in rows if r[col - 1] is not None])
            ws.column_dimensions[letter].width = min(width + 2, 40)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_report(path, items, config_models, judges, answers, judgements):
    """Write every saved result to the Excel file: all models ever evaluated, all cases in the dataset.

    Only judgements by the configured judges, made on the latest version of each answer, are counted.
    An existing file is updated in place, so sheets you added yourself are kept.
    """
    items_by_key = {(i["case_id"], i["question_key"]): i for i in items}
    answers = [a for a in answers.records.values() if (a["case_id"], a["question_key"]) in items_by_key]
    answers_by_key = {a["key"]: a for a in answers}
    judged = [j for j in judgements.records.values()
              if j["judge"] in judges and j["answer_key"] in answers_by_key and is_current(j, answers_by_key)]
    evaluated = {a["model"] for a in answers}
    models = [m for m in config_models if m in evaluated] + sorted(evaluated - set(config_models))

    def answer_of(j):
        return answers_by_key[j["answer_key"]]

    def item_of(a):
        return items_by_key[(a["case_id"], a["question_key"])]

    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    # Summary: one row per model, scores averaged over cases, runs and judges
    rows = []
    for model in models:
        model_answers = [a for a in answers if a["model"] == model]
        model_judged = [j for j in judged if answer_of(j)["model"] == model]
        scores = {d: avg([j["judgement"]["scores"][d]["score"] for j in model_judged]) for d in DIMENSIONS}
        cost = sum(a["cost_usd"] or 0 for a in model_answers)
        rows.append([model, len(model_answers), len(model_judged)] + [scores[d] for d in DIMENSIONS] +
                    [avg([a["words"] for a in model_answers]), avg([a["latency_s"] for a in model_answers]),
                     round(cost, 4)])
    add_sheet(wb, "Summary", ["Model", "Answers", "Judgements"] + DIMENSIONS +
              ["Avg words", "Avg latency (s)", "Answer cost ($)"], rows)

    # Scores by difficulty
    rows = []
    for model in models:
        for difficulty in sorted({i["difficulty"] for i in items}):
            js = [j for j in judged if answer_of(j)["model"] == model and item_of(answer_of(j))["difficulty"] == difficulty]
            if js:
                rows.append([model, difficulty, len(js)] +
                            [avg([j["judgement"]["scores"][d]["score"] for j in js]) for d in DIMENSIONS])
    add_sheet(wb, "By difficulty", ["Model", "Difficulty", "Judgements"] + DIMENSIONS, rows)

    # Scores: one row per answer and judge
    rows = []
    for j in sorted(judged, key=lambda j: j["key"]):
        a = answer_of(j)
        item = item_of(a)
        data = j["judgement"]
        kp = data.get("key_points") or []
        covered = {c: sum(1 for p in kp if p.get("covered") == c) for c in ("full", "partial", "none")}
        rows.append([a["case_id"], item["case_name"], item["difficulty"], a["question_key"], a["model"], a["run"],
                     j["judge"]] + [data["scores"][d]["score"] for d in DIMENSIONS] +
                    [covered["full"], covered["partial"], covered["none"],
                     data.get("strengths", ""), data.get("weaknesses", "")])
    add_sheet(wb, "Scores", ["Case", "Case name", "Difficulty", "Question", "Model", "Run", "Judge"] + DIMENSIONS +
              ["Key points full", "Key points partial", "Key points missed", "Strengths", "Weaknesses"],
              rows, wrap=("Strengths", "Weaknesses"))

    # Judge reasoning per dimension, plus the key point checklist
    rows = []
    for j in sorted(judged, key=lambda j: j["key"]):
        a = answer_of(j)
        data = j["judgement"]
        checklist = "\n".join(f"[{p.get('covered')}] {p.get('point')}" for p in data.get("key_points") or [])
        rows.append([a["case_id"], a["question_key"], a["model"], a["run"], j["judge"], checklist] +
                    [data["scores"][d]["reasoning"] for d in DIMENSIONS])
    add_sheet(wb, "Judge reasoning", ["Case", "Question", "Model", "Run", "Judge", "Key points"] + DIMENSIONS,
              rows, wrap=["Key points"] + DIMENSIONS)

    # Raw answers
    rows = [[a["case_id"], item_of(a)["case_name"], a["question_key"], a["model"], a["run"], a["words"],
             a["latency_s"], a["prompt_tokens"], a["completion_tokens"], a["cost_usd"], a["text"]]
            for a in sorted(answers, key=lambda a: a["key"])]
    add_sheet(wb, "Answers", ["Case", "Case name", "Question", "Model", "Run", "Words", "Latency (s)",
                              "Prompt tokens", "Completion tokens", "Cost ($)", "Answer"], rows, wrap=("Answer",))

    # Judge agreement: how often judges give the same score to the same answer
    by_answer = {}
    for j in judged:
        by_answer.setdefault(j["answer_key"], []).append(j)
    pairs = [js for js in by_answer.values() if len(js) > 1]
    rows = []
    for d in DIMENSIONS:
        diffs = []
        for js in pairs:
            s = [j["judgement"]["scores"][d]["score"] for j in js]
            diffs += [abs(x - y) for i, x in enumerate(s) for y in s[i + 1:]]
        rows.append([d, len(diffs), avg(diffs),
                     round(100 * sum(1 for x in diffs if x == 0) / len(diffs), 1) if diffs else None,
                     round(100 * sum(1 for x in diffs if x <= 1) / len(diffs), 1) if diffs else None])
    add_sheet(wb, "Judge agreement", ["Dimension", "Judge pairs", "Mean abs difference", "% exact match",
                                      "% within 1 point"], rows)

    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"\nCould not save {path}: close it in Excel, then run: python -m eval.run_eval --report-only\n"
                 "(All results are already saved in the results/*.jsonl files.)")
    judge_cost = sum(j["cost_usd"] or 0 for j in judged)
    answer_cost = sum(a["cost_usd"] or 0 for a in answers)
    print(f"\nSaved {path} ({len(answers)} answers, {len(judged)} judgements, "
          f"cost ${answer_cost:.2f} answers + ${judge_cost:.2f} judging)")
