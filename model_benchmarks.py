"""Export LLM benchmark scores, pricing and speed to a table and an Excel file.

Data comes from the Artificial Analysis API (free key: https://artificialanalysis.ai/documentation).
Set the key in the AA_API_KEY environment variable. Responses are cached for a day in .cache/
because the API allows 1,000 requests per day.

Usage:
    python model_benchmarks.py                          # all benchmarks -> model_benchmarks.xlsx
    python model_benchmarks.py --list                   # show available benchmarks
    python model_benchmarks.py --benchmarks gpqa,hle,mmlu_pro --sort gpqa
    python model_benchmarks.py --creator anthropic,openai,google --output shortlist.xlsx
    python model_benchmarks.py --refresh                # ignore the cache
"""

import argparse
import json
import os
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

API_URL = "https://artificialanalysis.ai/api/v2/data/llms/models"
CACHE_FILE = Path(__file__).parent / ".cache" / "aa_models.json"
CACHE_MAX_AGE = 24 * 60 * 60
DEFAULT_OUTPUT = Path(__file__).parent / "model_benchmarks.xlsx"

# Columns after the benchmarks: (header, getter, number format)
EXTRA_COLUMNS = [
    ("Input cost ($/1M tokens)", lambda m: (m.get("pricing") or {}).get("price_1m_input_tokens"), "0.00"),
    ("Output cost ($/1M tokens)", lambda m: (m.get("pricing") or {}).get("price_1m_output_tokens"), "0.00"),
    ("Speed (output tokens/s)", lambda m: m.get("median_output_tokens_per_second"), "0.0"),
]


def load_models(refresh, from_file):
    if from_file:
        with open(from_file, encoding="utf-8") as f:
            return json.load(f)["data"]

    if not refresh and CACHE_FILE.exists() and time.time() - CACHE_FILE.stat().st_mtime < CACHE_MAX_AGE:
        with open(CACHE_FILE, encoding="utf-8") as f:
            return json.load(f)["data"]

    api_key = os.environ.get("AA_API_KEY")
    if not api_key:
        sys.exit("Set the AA_API_KEY environment variable (get a free key at https://artificialanalysis.ai).")

    request = urllib.request.Request(API_URL, headers={"x-api-key": api_key})
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            payload = json.load(response)
    except urllib.error.HTTPError as e:
        sys.exit(f"API request failed: HTTP {e.code} {e.reason}")

    CACHE_FILE.parent.mkdir(exist_ok=True)
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(payload, f)
    return payload["data"]


def available_benchmarks(models):
    return sorted({k for m in models for k in (m.get("evaluations") or {})})


def build_table(models, benchmarks):
    """Return (headers, number formats, rows); missing values are None."""
    headers = ["Model Name", "Creator"] + benchmarks + [c[0] for c in EXTRA_COLUMNS]
    formats = [None, None] + ["0.000"] * len(benchmarks) + [c[2] for c in EXTRA_COLUMNS]
    rows = []
    for m in models:
        evals = m.get("evaluations") or {}
        row = [m.get("name", ""), (m.get("model_creator") or {}).get("name", "")]
        row += [evals.get(b) for b in benchmarks]
        row += [getter(m) for _, getter, _ in EXTRA_COLUMNS]
        rows.append(row)
    return headers, formats, rows


def write_excel(path, headers, formats, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Models"
    ws.append(headers)
    for row in rows:
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        width = max([len(str(header))] + [len(str(r[col - 1])) for r in rows if r[col - 1] is not None])
        ws.column_dimensions[get_column_letter(col)].width = min(width + 2, 45)
        if formats[col - 1]:
            for r in range(2, len(rows) + 2):
                ws.cell(row=r, column=col).number_format = formats[col - 1]

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def fmt(value):
    if value is None:
        return "-"
    return f"{value:.3g}" if isinstance(value, float) else str(value)


def print_table(headers, rows, limit):
    widths = [max(len(h), *(len(fmt(r[i])) for r in rows[:limit])) if rows else len(h) for i, h in enumerate(headers)]
    widths[0] = min(widths[0], 40)
    print("  ".join(h[:w].ljust(w) if i < 2 else h[:w].rjust(w) for i, (h, w) in enumerate(zip(headers, widths))))
    print("-" * (sum(widths) + 2 * (len(widths) - 1)))
    for r in rows[:limit]:
        cells = [fmt(v) for v in r]
        print("  ".join(c[:w].ljust(w) if i < 2 else c.rjust(w) for i, (c, w) in enumerate(zip(cells, widths))))
    if len(rows) > limit:
        print(f"... {len(rows) - limit} more rows in the Excel file")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--benchmarks", help="comma-separated benchmarks to include (default: all)")
    parser.add_argument("--creator", help="comma-separated model creator slugs to keep, e.g. anthropic,openai")
    parser.add_argument("--sort", help="column to sort by, highest first (a benchmark name, 'input', 'output' or 'speed')")
    parser.add_argument("--top", type=int, default=20, help="rows to print in the terminal (default 20); Excel gets all")
    parser.add_argument("--output", default=DEFAULT_OUTPUT, help=f"Excel file to write (default {DEFAULT_OUTPUT.name})")
    parser.add_argument("--list", action="store_true", help="list available benchmarks and exit")
    parser.add_argument("--refresh", action="store_true", help="fetch fresh data instead of using the cache")
    parser.add_argument("--from-file", help="read models from a saved API response instead of the API")
    args = parser.parse_args()

    models = load_models(args.refresh, args.from_file)
    all_benchmarks = available_benchmarks(models)

    if args.list:
        print("Available benchmarks:")
        for name in all_benchmarks:
            have = sum(1 for m in models if (m.get("evaluations") or {}).get(name) is not None)
            print(f"  {name:<45}{have:>5} models")
        return

    if args.creator:
        keep = {c.strip().lower() for c in args.creator.split(",")}
        models = [m for m in models if (m.get("model_creator") or {}).get("slug", "").lower() in keep]

    benchmarks = [b.strip() for b in args.benchmarks.split(",")] if args.benchmarks else all_benchmarks
    unknown = set(benchmarks) - set(all_benchmarks)
    if unknown:
        sys.exit(f"Unknown benchmark(s): {', '.join(sorted(unknown))}. Run with --list to see options.")

    headers, formats, rows = build_table(models, benchmarks)

    if args.sort:
        aliases = {"input": EXTRA_COLUMNS[0][0], "output": EXTRA_COLUMNS[1][0], "speed": EXTRA_COLUMNS[2][0]}
        sort_header = aliases.get(args.sort, args.sort)
        if sort_header not in headers:
            sys.exit(f"Cannot sort by '{args.sort}': not a column in this table.")
        i = headers.index(sort_header)
        rows.sort(key=lambda r: (r[i] is not None, r[i] or 0), reverse=True)
    else:
        rows.sort(key=lambda r: r[0].lower())

    print_table(headers, rows, args.top)
    write_excel(args.output, headers, formats, rows)
    print(f"\nSaved {len(rows)} models to {args.output}")


if __name__ == "__main__":
    main()
